from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

from app.ingestion import DatasetParseError, parse_dataset


def test_csv_preserves_values_and_source_rows(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("nama,jumlah,catatan\nANI,0,N/A\nBUDI,2,   \n", encoding="utf-8")

    parsed = parse_dataset(path, "data.csv")

    assert [column.original_name for column in parsed.columns] == ["nama", "jumlah", "catatan"]
    assert parsed.source_rows == [2, 3]
    assert parsed.records == [
        {"nama": "ANI", "jumlah": "0", "catatan": "N/A"},
        {"nama": "BUDI", "jumlah": "2", "catatan": "   "},
    ]


def test_csv_supports_windows_1252_and_semicolon(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_bytes("wilayah;keterangan\nJAKARTA;Café\n".encode("cp1252"))

    parsed = parse_dataset(path, "data.csv")

    assert parsed.records == [{"wilayah": "JAKARTA", "keterangan": "Café"}]


def test_xlsx_uses_first_nonempty_sheet_and_real_row_numbers(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.title = "Kosong"
    sheet = workbook.create_sheet("Data OPD")
    sheet.append([])
    sheet.append([])
    sheet.append(["nama", "jumlah", "catatan"])
    sheet.append(["ANI", 0, "N/A"])
    workbook.save(path)

    parsed = parse_dataset(path, "data.xlsx")

    assert parsed.worksheet_name == "Data OPD"
    assert parsed.source_rows == [4]
    assert parsed.records == [{"nama": "ANI", "jumlah": 0, "catatan": "N/A"}]


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("nama,nama\nANI,ANI\n", "duplikat"),
        ("", "kosong"),
        ("nama\n", "minimal satu baris"),
    ],
)
def test_csv_rejects_ambiguous_or_empty_tables(
    tmp_path: Path, content: str, message: str
) -> None:
    path = tmp_path / "data.csv"
    path.write_text(content, encoding="utf-8")

    with pytest.raises(DatasetParseError, match=message):
        parse_dataset(path, "data.csv")


def test_xlsx_rejects_invalid_archive(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    path.write_bytes(b"not-an-xlsx")

    with pytest.raises(DatasetParseError, match="tidak valid"):
        parse_dataset(path, "data.xlsx")


def test_csv_ignores_physical_blank_lines_and_keeps_real_row_number(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("nama,jumlah\n\nANI,0\n\n", encoding="utf-8")

    parsed = parse_dataset(path, "data.csv")

    assert parsed.source_rows == [3]
    assert parsed.records == [{"nama": "ANI", "jumlah": "0"}]


def test_csv_preserves_newlines_inside_quoted_values(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text('nama,alamat\nANI,"Jalan\nMawar"\n', encoding="utf-8")

    parsed = parse_dataset(path, "data.csv")

    assert parsed.source_rows == [2]
    assert parsed.records == [{"nama": "ANI", "alamat": "Jalan\nMawar"}]


def test_csv_supports_a_single_column_dataset(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("nama\nANI\nBUDI\n", encoding="utf-8")

    parsed = parse_dataset(path, "data.csv")

    assert parsed.source_rows == [2, 3]
    assert parsed.records == [{"nama": "ANI"}, {"nama": "BUDI"}]


@pytest.mark.parametrize("headers", [["Kolom 2", ""], ["_source_row", "nama"]])
def test_rejects_ambiguous_or_reserved_effective_headers(
    tmp_path: Path, headers: list[str]
) -> None:
    path = tmp_path / "data.csv"
    path.write_text(",".join(headers) + "\n1,2\n", encoding="utf-8")

    with pytest.raises(DatasetParseError, match="ambigu|_source_row"):
        parse_dataset(path, "data.csv")


@pytest.mark.parametrize(
    ("content", "limits", "message"),
    [
        ("a,b,c\n1,2,3\n", {"max_columns": 2}, "kolom"),
        ("a\n1\n2\n", {"max_rows": 1}, "baris"),
    ],
)
def test_csv_enforces_dataset_bounds(
    tmp_path: Path, content: str, limits: dict[str, int], message: str
) -> None:
    path = tmp_path / "data.csv"
    path.write_text(content, encoding="utf-8")

    with pytest.raises(DatasetParseError, match=message):
        parse_dataset(path, "data.csv", **limits)


def test_xlsx_enforces_dataset_row_bound(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.append(["nama"])
    workbook.active.append(["ANI"])
    workbook.active.append(["BUDI"])
    workbook.save(path)

    with pytest.raises(DatasetParseError, match="baris"):
        parse_dataset(path, "data.xlsx", max_rows=1)


def test_xlsx_rejects_dimensions_over_cell_bound_before_iteration(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.append(["a", "b"])
    workbook.active.append([1, 2])
    workbook.save(path)

    with pytest.raises(DatasetParseError, match="sel"):
        parse_dataset(path, "data.xlsx", max_cells=3)


def test_rejects_effective_header_over_255_characters(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text(f"{'a' * 256}\nvalue\n", encoding="utf-8")

    with pytest.raises(DatasetParseError, match="255"):
        parse_dataset(path, "data.csv")


def test_xlsx_skips_styled_blank_rows_and_preserves_source_rows(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["nama"])
    sheet.append(["ANI"])
    sheet[4][0].number_format = "0"
    workbook.save(path)

    parsed = parse_dataset(path, "data.xlsx")

    assert parsed.records == [{"nama": "ANI"}]
    assert parsed.source_rows == [2]


def test_xlsx_rejects_header_with_only_blank_data_rows(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["nama"])
    sheet[3][0].number_format = "0"
    workbook.save(path)

    with pytest.raises(DatasetParseError, match="minimal satu baris data"):
        parse_dataset(path, "data.xlsx")


def test_xlsx_preserves_formula_only_row_without_cached_result(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.append(["hasil"])
    workbook.active.append(["=1+1"])
    workbook.save(path)

    parsed = parse_dataset(path, "data.xlsx")

    assert parsed.records == [{"hasil": None}]
    assert parsed.source_rows == [2]


def test_xlsx_preserves_array_formula_only_row(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.append(["hasil"])
    workbook.active["A2"] = ArrayFormula("A2", "=1+1")
    workbook.save(path)

    parsed = parse_dataset(path, "data.xlsx")

    assert parsed.records == [{"hasil": None}]
