from io import BytesIO
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import CheckConstraint, inspect

import app.datasets as datasets_module
from app.config import Settings, get_settings
from app.main import app
from app.models import Dataset
from app.quality import QualityAssessmentError


def test_config_exposes_upload_limit(client: TestClient) -> None:
    response = client.get("/api/config")

    assert response.status_code == 200
    assert response.json() == {"max_upload_mb": 1, "backoffice_enabled": False, "llm_max_candidates": 50}


def test_assessment_verdict_values_are_constrained_by_schema() -> None:
    def constraints(model: type[object]) -> str:
        return "\n".join(
            str(constraint.sqltext)
            for constraint in inspect(model).local_table.constraints
            if isinstance(constraint, CheckConstraint)
        )

    assert "PASSED" in constraints(Dataset)
    assert "NEEDS_CORRECTION" in constraints(Dataset)


def test_upload_list_and_detail_dataset(client: TestClient) -> None:
    response = client.post(
        "/api/datasets",
        files={"file": ("contoh.csv", b"nama,jumlah\nANI,0\nBUDI,2\n", "text/csv")},
    )

    assert response.status_code == 201
    uploaded = response.json()
    assert uploaded["status"] == "COMPLETED"
    assert (uploaded["row_count"], uploaded["column_count"]) == (2, 2)
    assert uploaded["finding_count"] == 0
    assert uploaded["findings"] == []
    assert "original_file_key" not in uploaded

    history = client.get("/api/datasets").json()
    assert [item["id"] for item in history] == [uploaded["id"]]

    detail = client.get(f"/api/datasets/{uploaded['id']}")
    assert detail.status_code == 200
    assert detail.json()["preview"][0] == {"_source_row": 2, "nama": "ANI", "jumlah": "0"}
    assert detail.json()["column_profiles"][1]["inferred_type"] == "integer"


def test_finding_context_and_preview_locate_sparse_excel_source_row(client: TestClient) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["nama_kegiatan", "jumlah"])
    sheet.append(["Kegiatan awal", 1])
    sheet.cell(40, 1, "Kegiatan akhir")
    sheet.cell(40, 2, 2)
    content = BytesIO()
    workbook.save(content)
    uploaded = client.post("/api/datasets", files={"file": ("sparse.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert uploaded.status_code == 201
    dataset_id = uploaded.json()["id"]
    findings = client.get(f"/api/datasets/{dataset_id}/findings?group_by_row=true&issue_type=INVALID_CASE").json()
    final_group = next(group for group in findings["row_groups"] if group["row"] == 40)
    assert final_group["source_record"]["nama_kegiatan"] == "Kegiatan akhir"
    assert final_group["source_record"]["_source_row"] == 40
    located = client.get(f"/api/datasets/{dataset_id}/preview?source_row=40&page_size=1")
    assert located.status_code == 200
    assert located.json()["page"] == 2
    assert located.json()["items"][0]["_source_row"] == 40
    assert client.get(f"/api/datasets/{dataset_id}/preview?source_row=39").status_code == 404


def test_grouped_findings_search_source_context_and_filter_kind(client: TestClient) -> None:
    uploaded = client.post("/api/datasets", files={"file": ("context.csv", b"lokasi,nama_kegiatan,jumlah\nGAMBIR,kegiatan satu,1\nCILINCING,kegiatan dua,2\n", "text/csv")})
    assert uploaded.status_code == 201
    dataset_id = uploaded.json()["id"]
    url = f"/api/datasets/{dataset_id}/findings"
    found = client.get(url, params={"group_by_row": True, "search": "CILINCING", "assessment_kind": "required"})
    assert found.status_code == 200
    body = found.json()
    assert body["total_rows"] == 1
    assert body["row_groups"][0]["row"] == 3
    assert body["row_groups"][0]["source_record"]["lokasi"] == "CILINCING"
    assert all(finding["issue_type"] == "INVALID_CASE" for finding in body["items"])
    empty = client.get(url, params={"group_by_row": True, "assessment_kind": "review"}).json()
    assert empty["total_rows"] == 0
    assert empty["row_groups"] == []
    assert client.get(url, params={"assessment_kind": "unsupported"}).status_code == 422


def test_grouped_finding_sort_is_global_numeric_stable_and_missing_last(client: TestClient) -> None:
    uploaded = client.post("/api/datasets", files={"file": ("sort.csv", b"nama,jumlah\nZulu,10\nAlpha,2\nMike,30\nDelta,\nEcho,0\nBravo,2\n", "text/csv")})
    assert uploaded.status_code == 201
    url = f"/api/datasets/{uploaded.json()['id']}/findings"
    params = {"group_by_row": True, "sort_by": "jumlah", "page_size": 2}
    ascending = [client.get(url, params={**params, "page": page}).json() for page in (1, 2, 3)]
    assert [group["row"] for result in ascending for group in result["row_groups"]] == [6, 3, 7, 2, 4, 5]
    descending = client.get(url, params={**params, "sort_direction": "desc", "page_size": 100}).json()
    assert [group["row"] for group in descending["row_groups"]] == [4, 2, 3, 7, 6, 5]
    text_sort = client.get(url, params={"group_by_row": True, "sort_by": "nama", "page_size": 100}).json()
    assert [group["row"] for group in text_sort["row_groups"]] == [3, 7, 5, 6, 4, 2]
    source_sort = client.get(url, params={"group_by_row": True, "sort_by": "_source_row", "sort_direction": "desc", "page_size": 1}).json()
    assert source_sort["row_groups"][0]["row"] == 7
    assert client.get(url, params={"group_by_row": True, "sort_by": "not_a_column"}).status_code == 422
    assert client.get(url, params={"sort_direction": "invalid"}).status_code == 422


def test_preview_profiles_file_without_creating_assessment(client: TestClient) -> None:
    response = client.post(
        "/api/datasets/preview",
        files={
            "file": (
                "preview.csv",
                b"nama,jumlah,tanggal_kejadian\nANI,10,20250101\nBUDI,20,20250102\n",
                "text/csv",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["filename"] == "preview.csv"
    assert (body["row_count"], body["column_count"]) == (2, 3)
    assert body["rows"][0] == {
        "_source_row": 2,
        "nama": "ANI",
        "jumlah": "10",
        "tanggal_kejadian": "20250101",
    }
    assert [column["candidate_role"] for column in body["columns"]] == [
        "text",
        "numeric",
        "date",
    ]
    assert body["checks"] == [
        "FULL_ROW_DUPLICATE",
        "MISSING_VALUE",
        "UPPERCASE",
        "DATE_FORMAT",
        "NUMERIC_FORMAT",
    ]
    assert client.get("/api/datasets").json() == []


def test_backoffice_failure_persists_only_safe_message(client: TestClient) -> None:
    app.dependency_overrides[get_settings] = lambda: Settings(
        database_url="sqlite://",
        backoffice_api_base_url="https://backoffice.test",
        backoffice_api_token="secret-token",
    )

    response = client.post("/api/backoffice-assessments")

    assert response.status_code == 422
    failed = client.get("/api/datasets").json()[0]
    assert failed["failure_message"] == "Backoffice assessment gagal."
    assert "secret-token" not in str(failed)


def test_findings_are_read_only_and_reassessment_can_be_compared(client: TestClient) -> None:
    first = client.post(
        "/api/datasets",
        files={"file": ("first.csv", b"nama\nAni\n", "text/csv")},
    ).json()
    finding_id = first["findings"][0]["id"]
    assert "status" not in first["findings"][0]

    changed = client.patch(f"/api/findings/{finding_id}", json={"status": "DISMISSED"})
    assert changed.status_code == 404
    assert client.get("/api/dashboard").json()["needs_correction_assessments"] == 1

    second = client.post(
        "/api/datasets",
        data={"parent_dataset_id": first["id"]},
        files={"file": ("second.csv", b"nama\nANI\n", "text/csv")},
    ).json()
    comparison = client.get(f"/api/datasets/{second['id']}/comparison").json()
    assert comparison["available"] is True
    assert comparison["resolved"] == [finding_id]


def test_reports_include_summary_detail_and_pdf(client: TestClient) -> None:
    dataset = client.post(
        "/api/datasets",
        files={"file": ("laporan.csv", b"nama\nAni\n", "text/csv")},
    ).json()

    xlsx = client.get(f"/api/datasets/{dataset['id']}/report?format=xlsx")
    pdf = client.get(f"/api/datasets/{dataset['id']}/report?format=pdf")

    workbook = __import__("openpyxl").load_workbook(BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Ringkasan", "Temuan"]
    assert workbook["Ringkasan"]["A1"].value == "Ringkasan assessment"
    assert workbook["Temuan"]["A1"].value == "Baris"
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF-")
    assert len(pdf.content) > 2000


def test_dashboard_reports_findings_from_latest_completed_assessment_only(
    client: TestClient,
) -> None:
    client.post(
        "/api/datasets",
        files={"file": ("first.csv", b"nama\nAni\n", "text/csv")},
    )
    latest = client.post(
        "/api/datasets",
        files={"file": ("latest.csv", b"catatan\n-\n", "text/csv")},
    ).json()

    dashboard = client.get("/api/dashboard").json()

    assert dashboard["latest_assessment"] == {
        "id": latest["id"],
        "original_filename": "latest.csv",
        "source_type": "file_upload",
        "completed_at": latest["completed_at"],
        "finding_count": 1,
        "issue_counts": {"INVALID_PLACEHOLDER": 1},
    }


def test_findings_are_paginated_and_keep_source_row_order(client: TestClient) -> None:
    dataset = client.post(
        "/api/datasets",
        files={
            "file": (
                "ordered.csv",
                b"nama,catatan\nANI,N/A\nAni,-\n,VALID\n",
                "text/csv",
            )
        },
    ).json()

    first_page = client.get(
        f"/api/datasets/{dataset['id']}/findings?page=1&page_size=2"
    ).json()
    second_page = client.get(
        f"/api/datasets/{dataset['id']}/findings?page=2&page_size=2"
    ).json()
    medium = client.get(
        f"/api/datasets/{dataset['id']}/findings?severity=MEDIUM&page_size=50"
    ).json()
    high = client.get(
        f"/api/datasets/{dataset['id']}/findings?severity=HIGH&page_size=50"
    ).json()
    first_row_page = client.get(
        f"/api/datasets/{dataset['id']}/findings?group_by_row=true&page=1&page_size=1"
    ).json()
    second_row_page = client.get(
        f"/api/datasets/{dataset['id']}/findings?group_by_row=true&page=2&page_size=1"
    ).json()
    invalid_case_rows = client.get(
        f"/api/datasets/{dataset['id']}/findings?severity=HIGH&issue_type=INVALID_CASE&group_by_row=true&page_size=10"
    ).json()

    assert first_page["total"] == 3
    assert first_page["page"] == 1
    assert first_page["page_size"] == 2
    assert first_page["total_rows"] == 2
    assert first_page["row_groups"] == []
    assert first_page["by_issue_type"] == {
        "INVALID_CASE": 1,
        "INVALID_PLACEHOLDER": 1,
        "MISSING_VALUE": 1,
    }
    assert [(item["row"], item["column"]) for item in first_page["items"]] == [
        (3, "nama"),
        (3, "catatan"),
    ]
    assert [(item["row"], item["column"]) for item in second_page["items"]] == [
        (4, "nama")
    ]
    assert {item["issue_type"] for item in medium["items"]} == {"INVALID_PLACEHOLDER"}
    assert medium["by_issue_type"] == {"INVALID_PLACEHOLDER": 1}
    assert {item["issue_type"] for item in high["items"]} == {
        "INVALID_CASE",
        "MISSING_VALUE",
    }
    assert high["by_issue_type"] == {"INVALID_CASE": 1, "MISSING_VALUE": 1}
    assert first_row_page["total"] == 3
    assert first_row_page["total_rows"] == 2
    assert first_row_page["row_groups"][0]["row"] == 3
    assert first_row_page["row_groups"][0]["action_summary_source"] == "fallback"
    assert first_row_page["row_groups"][0]["action_summary"].startswith("Memeriksa")
    assert first_row_page["row_groups"][0]["action_summary_confidence"] is None
    assert [item["column"] for item in first_row_page["row_groups"][0]["findings"]] == [
        "nama",
        "catatan",
    ]
    assert [item["row"] for item in first_row_page["items"]] == [3, 3]
    assert second_row_page["row_groups"][0]["row"] == 4
    assert [item["column"] for item in second_row_page["row_groups"][0]["findings"]] == [
        "nama"
    ]
    assert invalid_case_rows["row_groups"][0]["action_summary_source"] == "fallback"
    assert invalid_case_rows["row_groups"][0]["action_summary"].startswith(
        "Mengganti nilai"
    )


def test_global_findings_are_read_only_searchable_and_include_dataset(client: TestClient) -> None:
    dataset = client.post(
        "/api/datasets",
        files={"file": ("lintas.csv", b"nama\nAni\n", "text/csv")},
    ).json()

    response = client.get("/api/findings?search=lintas&page_size=10")

    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 1
    assert body["items"][0]["dataset_id"] == dataset["id"]
    assert body["items"][0]["dataset_name"] == "lintas.csv"
    assert body["items"][0]["row"] == 2


def test_reference_data_and_safe_settings_are_api_backed(client: TestClient) -> None:
    reference = client.get("/api/reference-data?search=SUNGAI%20BAMBU&page_size=10")
    settings = client.get("/api/settings")

    assert reference.status_code == 200
    assert reference.json()["total"] == 1
    assert reference.json()["items"][0]["village_name"] == "SUNGAI BAMBU"
    assert reference.json()["summary"]["villages"] == 267
    assert settings.status_code == 200
    assert settings.json()["limits"]["upload_mb"] == 1
    assert settings.json()["semantic"]["enabled"] is False
    assert settings.json()["formats"]["upload"] == ["CSV", "XLSX", "XLS"]
    assert settings.json()["formats"]["report"] == ["CSV", "XLSX", "PDF"]
    assert "api_key" not in str(settings.json()).lower()


def test_upload_runs_structural_quality_checks(client: TestClient) -> None:
    response = client.post(
        "/api/datasets",
        files={
            "file": (
                "quality.csv",
                b"nama,catatan\nANI,N/A\nAni,-\n,0\nANI,N/A\n",
                "text/csv",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["finding_count"] == 4
    execution = body["source_metadata"]["quality_checks"]
    assert set(execution["executed"]) == set(execution["relevant"])
    assert "DATE_FORMAT" not in execution["executed"]
    assert "DATE_FORMAT" in execution["available"]
    assert body["finding_summary"]["by_issue_type"] == {
        "DUPLICATE": 1,
        "MISSING_VALUE": 1,
        "INVALID_PLACEHOLDER": 1,
        "INVALID_CASE": 1,
    }
    assert body["finding_summary"]["by_severity"] == {"HIGH": 3, "MEDIUM": 1}
    assert [(item["row"], item["issue_type"]) for item in body["findings"]] == [
        (3, "INVALID_CASE"),
        (3, "INVALID_PLACEHOLDER"),
        (4, "MISSING_VALUE"),
        (5, "DUPLICATE"),
    ]
    missing = next(item for item in body["findings"] if item["issue_type"] == "MISSING_VALUE")
    assert missing["expected_value"] == "N/A"
    assert missing["recommendation"] == 'Mengganti nilai kosong dengan "N/A" sebagai representasi nilai yang tidak tersedia.'
    assert "original_file_key" not in str(body)
    assert client.get("/api/datasets").json()[0]["finding_count"] == 4


def test_upload_xlsx_runs_structural_quality_checks(client: TestClient) -> None:
    workbook = Workbook()
    workbook.active.append(["nama"])
    workbook.active.append(["Ani"])
    content = BytesIO()
    workbook.save(content)

    response = client.post(
        "/api/datasets",
        files={"file": ("quality.xlsx", content.getvalue(), "application/octet-stream")},
    )

    assert response.status_code == 201
    assert response.json()["findings"][0]["issue_type"] == "INVALID_CASE"


def test_upload_runs_administrative_reference_checks(client: TestClient) -> None:
    response = client.post(
        "/api/datasets",
        files={
            "file": (
                "reference.csv",
                b"wilayah,kecamatan,kelurahan\nBUKAN WILAYAH,TANJUNG PRIOK,SUNGAI BAMBU\n",
                "text/csv",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["finding_summary"]["by_issue_type"] == {"INVALID_REFERENCE": 1}
    assert body["column_profiles"][0]["candidate_role"] == "administrative_city"


def test_xlsx_native_administrative_value_is_json_safe(client: TestClient) -> None:
    workbook = Workbook()
    workbook.active.append(["wilayah"])
    workbook.active.append([date(2025, 1, 29)])
    content = BytesIO()
    workbook.save(content)

    response = client.post(
        "/api/datasets",
        files={"file": ("reference.xlsx", content.getvalue(), "application/octet-stream")},
    )

    assert response.status_code == 201
    assert response.json()["findings"][0]["original_value"] == "2025-01-29T00:00:00"


def test_invalid_upload_is_rejected_without_internal_paths(client: TestClient) -> None:
    response = client.post(
        "/api/datasets",
        files={"file": ("rusak.xlsx", b"not-xlsx", "application/octet-stream")},
    )

    assert response.status_code == 422
    body = response.json()
    assert "tidak valid" in body["detail"]["message"]
    assert body["detail"]["request_id"] == response.headers["x-request-id"]
    assert "storage" not in str(body).lower()


def test_upload_xlsx_uses_the_real_file_type_while_staged(client: TestClient) -> None:
    workbook = Workbook()
    workbook.active.append(["nama", "jumlah"])
    workbook.active.append(["ANI", 0])
    content = BytesIO()
    workbook.save(content)

    response = client.post(
        "/api/datasets",
        files={"file": ("contoh.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 201
    assert response.json()["worksheet_name"] == "Sheet"


def test_upload_legacy_xls_uses_real_biff_fixture(client: TestClient) -> None:
    content = (Path(__file__).parent / "fixtures" / "legacy_sample.xls").read_bytes()

    response = client.post(
        "/api/datasets",
        files={"file": ("legacy_sample.xls", content, "application/vnd.ms-excel")},
    )

    assert response.status_code == 201
    body = response.json()
    assert body["file_type"] == "xls"
    assert body["worksheet_name"] == "DATA LAMA"
    assert body["row_count"] == 2
    assert body["preview"][0]["tanggal"] == "2026-09-01T00:00:00"


def test_upload_xlsx_serializes_native_date_values(client: TestClient) -> None:
    workbook = Workbook()
    workbook.active.append(["tanggal_kejadian"])
    workbook.active.append([date(2025, 1, 29)])
    content = BytesIO()
    workbook.save(content)

    response = client.post(
        "/api/datasets",
        files={"file": ("tanggal.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 201
    body = response.json()
    assert body["preview"][0]["tanggal_kejadian"] == "2025-01-29T00:00:00"
    assert body["column_profiles"][0]["sample_values"] == ["2025-01-29T00:00:00"]
    assert body["findings"][0]["original_value"] == "2025-01-29T00:00:00"
    assert body["findings"][0]["expected_value"] == "20250129"


def test_upload_persists_date_and_period_findings(client: TestClient) -> None:
    response = client.post(
        "/api/datasets",
        files={
            "file": (
                "tanggal.csv",
                b"tanggal_kejadian,periode_data\n20250129,202501\n2025-01-30,202513\n",
                "text/csv",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["finding_summary"]["by_issue_type"] == {"INVALID_DATE_FORMAT": 2}
    assert body["finding_summary"]["by_severity"] == {"HIGH": 2}
    assert [item["expected_value"] for item in body["findings"]] == [
        "20250130",
        "YYYY, YYYYMM, YYYY-T1–T4 atau YYYY-S1–S2",
    ]
    assert all(item["detector"] == "deterministic:date-v1" for item in body["findings"])


def test_upload_persists_numeric_format_findings(client: TestClient) -> None:
    response = client.post(
        "/api/datasets",
        files={
            "file": (
                "numeric.csv",
                b'ukuran\n0\n1\n2\n3\n4\n5\n6\n7\n8\n"1,5"\n',
                "text/csv",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["finding_summary"]["by_issue_type"] == {"FORMAT_ERROR": 1}
    finding = body["findings"][0]
    assert (finding["row"], finding["column"]) == (11, "ukuran")
    assert finding["original_value"] == "1,5"
    assert finding["expected_value"] == "Angka ASCII dengan titik desimal"
    assert finding["detector"] == "deterministic:numeric-v1"


def test_upload_xlsx_preserves_native_numbers_for_numeric_quality(
    client: TestClient,
) -> None:
    workbook = Workbook()
    workbook.active.append(["ukuran"])
    for value in range(9):
        workbook.active.append([value])
    workbook.active.append(["1,5"])
    content = BytesIO()
    workbook.save(content)

    response = client.post(
        "/api/datasets",
        files={"file": ("numeric.xlsx", content.getvalue(), "application/octet-stream")},
    )

    assert response.status_code == 201
    finding = response.json()["findings"][0]
    assert finding["row"] == 11
    assert finding["original_value"] == "1,5"


def test_unexpected_processing_failure_is_safe_and_cleans_artifacts(
    client: TestClient, tmp_path: Path, monkeypatch
) -> None:
    def fail_snapshot(*_args, **_kwargs):
        raise RuntimeError("internal detail must stay private")

    monkeypatch.setattr(datasets_module, "write_snapshot", fail_snapshot)

    response = client.post(
        "/api/datasets",
        files={"file": ("contoh.csv", b"nama,jumlah\nANI,0\n", "text/csv")},
    )

    assert response.status_code == 500
    assert response.json()["detail"]["message"] == "Dataset tidak dapat diproses."
    assert client.get("/api/datasets").json()[0]["status"] == "FAILED"
    assert list((tmp_path / "storage" / "originals").glob("*")) == []


def test_quality_failure_is_safe_and_cleans_artifacts(
    client: TestClient, tmp_path: Path, monkeypatch
) -> None:
    def fail_quality(*_args, **_kwargs):
        raise RuntimeError("quality detail must stay private")

    monkeypatch.setattr(datasets_module, "quality_findings", fail_quality)

    response = client.post(
        "/api/datasets",
        files={"file": ("contoh.csv", b"nama\nANI\n", "text/csv")},
    )

    assert response.status_code == 500
    assert response.json()["detail"]["message"] == "Dataset tidak dapat diproses."
    assert client.get("/api/datasets").json()[0]["status"] == "FAILED"
    assert list((tmp_path / "storage" / "originals").glob("*")) == []
    assert list((tmp_path / "storage" / "snapshots").glob("*")) == []


def test_quality_limit_returns_actionable_422_and_cleans_artifacts(
    client: TestClient, tmp_path: Path, monkeypatch
) -> None:
    def reject_findings(_parsed, _profile, _reference, max_findings, *, executed_checks=None):
        assert max_findings == 10_000
        raise QualityAssessmentError(
            "Dataset menghasilkan lebih dari 10000 temuan kualitas. Kurangi ukuran dataset."
        )

    monkeypatch.setattr(datasets_module, "quality_findings", reject_findings)

    response = client.post(
        "/api/datasets",
        files={"file": ("contoh.csv", b"nama\nAni\n", "text/csv")},
    )

    assert response.status_code == 422
    assert "kurangi ukuran dataset" in response.json()["detail"]["message"].lower()
    assert client.get("/api/datasets").json()[0]["status"] == "FAILED"
    assert list((tmp_path / "storage" / "originals").glob("*")) == []
    assert list((tmp_path / "storage" / "snapshots").glob("*")) == []


def test_reference_failure_returns_503_and_cleans_artifacts(
    client: TestClient, tmp_path: Path
) -> None:
    app.dependency_overrides[get_settings] = lambda: Settings(
        database_url="sqlite://",
        storage_root=tmp_path / "storage",
        master_reference_path=tmp_path / "missing.csv",
    )

    response = client.post(
        "/api/datasets",
        files={"file": ("contoh.csv", b"wilayah\nJAKARTA\n", "text/csv")},
    )

    assert response.status_code == 503
    assert "referensi" in response.json()["detail"]["message"].lower()
    assert client.get("/api/datasets").json()[0]["status"] == "FAILED"
    assert list((tmp_path / "storage" / "originals").glob("*")) == []
    assert list((tmp_path / "storage" / "snapshots").glob("*")) == []


def test_upload_rejects_an_unsafe_filename_before_persistence(client: TestClient) -> None:
    response = client.post(
        "/api/datasets",
        files={"file": (f"{'a' * 256}.csv", b"nama\nANI\n", "text/csv")},
    )

    assert response.status_code == 422
    assert response.json()["detail"]["message"] == "Nama file terlalu panjang."
    assert client.get("/api/datasets").json() == []
