from collections import Counter
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import pytest
from openpyxl import Workbook

from app.ingestion import ParsedColumn, ParsedDataset, parse_dataset
from app.profiling import profile_dataset
from app.quality import (
    QualityAssessmentError, identifier_findings, planned_quality_checks,
    quality_findings, structural_findings,
)
from app.reference import load_administrative_reference
from app.semantic import apply_semantic_roles


def dataset(records):
    return ParsedDataset(
        [ParsedColumn(name, name, name) for name in records[0]],
        records, list(range(2, len(records) + 2)),
    )


@pytest.mark.parametrize("name", [
    "nomor_KK", "no. KK", "Nomor Kartu Keluarga", "NIK", "nomor_nik",
    "nomor_SK", "No.Izin", "kode_produk", "id_pelanggan", "transaksi_id", "rt", "rw",
])
def test_identifier_headers_are_not_numeric_measures(name):
    data = dataset([{name: str(1111111111111111 + i)} for i in range(8)])
    profile = profile_dataset(data)
    assert profile.columns[0].candidate_role == "identifier"
    checks = planned_quality_checks(profile, semantic_enabled=True)
    assert not set(checks) & {"STATISTICAL_ANOMALY", "CALCULATION", "NUMERIC_FORMAT", "SEMANTIC_TYPO"}
    assert apply_semantic_roles(profile, {name: "percentage"}).columns[0].candidate_role == "identifier"


def test_measure_columns_remain_numeric_and_select_iqr():
    data = dataset([
        {"jumlah": i, "usia": i + 20, "pendapatan": i * 1000}
        for i in range(1, 9)
    ])
    profile = profile_dataset(data)
    assert all(column.candidate_role == "numeric" for column in profile.columns)
    assert "STATISTICAL_ANOMALY" in planned_quality_checks(profile, semantic_enabled=False)


def test_identifier_format_preserves_source_and_skips_missing_markers():
    values = ["0123456789012345", 1234567890123456, "123", "1" * 17,
              "1a" + "1" * 14, True, 1234567890123456.0, " 1234567890123456",
              "1.234567890123456E15", None, "", " ", "-", "N/A", 0, "0"]
    data = dataset([{"nomor_kk": value} for value in values])
    findings = identifier_findings(data)
    assert [f.row for f in findings] == [4, 5, 6, 7, 8, 9, 10, 16, 17]
    assert all(f.expected_value is None for f in findings)
    assert findings[0].original_value == "123"
    assert data.records[0]["nomor_kk"] == "0123456789012345"
    with pytest.raises(QualityAssessmentError):
        identifier_findings(data, max_findings=1)


def test_repeated_household_identifiers_do_not_imply_duplicate_rows():
    data = dataset([
        {"nomor_kk": "0123456789012345", "nama": "ANI"},
        {"nomor_kk": "0123456789012345", "nama": "BUDI"},
    ])
    assert identifier_findings(data) == []
    assert structural_findings(data) == []


def test_document_identifier_has_no_invented_format_or_uppercase_rule():
    data = dataset([{"No.Izin": "1a.0312Bc", "nomor_SK": "1a.0312Bc", "nama": "Ani"}])
    assert identifier_findings(data) == []
    assert [(f.column, f.issue_type) for f in structural_findings(data)] == [("nama", "INVALID_CASE")]


def test_known_identifiers_do_not_run_numeric_detectors(monkeypatch):
    from app import quality
    data = dataset([{"nomor_kk": str(1111111111111111 + i), "nik": str(2222222222222222 + i), "id_pelanggan": i} for i in range(8)])
    def forbidden(*args, **kwargs):
        raise AssertionError("Identifier dispatched as measure")
    for name in ["numeric_findings", "statistical_findings", "calculation_findings"]:
        monkeypatch.setattr(quality, name, forbidden)
    reference = load_administrative_reference(Path(__file__).parents[1] / "reference/master.csv")
    executed = []
    assert quality_findings(data, profile_dataset(data), reference, executed_checks=executed) == []
    assert "IDENTIFIER_FORMAT" in executed


def test_styled_empty_column_removed_but_named_empty_column_preserved(tmp_path):
    path = tmp_path / "styled.xlsx"
    book = Workbook(); sheet = book.active
    sheet.append(["nama", "catatan"]); sheet.append(["ANI", None])
    sheet["C1"].number_format = "@"; sheet["C2"].number_format = "@"
    book.save(path); book.close()
    before = path.read_bytes()
    data = parse_dataset(path, path.name)
    assert [c.display_name for c in data.columns] == ["nama", "catatan"]
    assert [(f.column, f.issue_type) for f in structural_findings(data)] == [("catatan", "MISSING_VALUE")]
    assert path.read_bytes() == before


@pytest.mark.parametrize("value", ["tertinggal", '=LEFT(A3,1)', '=IF(A3="ANI","",A3)'])
def test_anonymous_column_with_data_or_formula_has_one_structural_finding(tmp_path, value):
    path = tmp_path / "anonymous.xlsx"
    book = Workbook(); sheet = book.active
    sheet.append(["nama", "nik", None])
    for index in range(8):
        sheet.append([f"ORANG {index}", str(1111111111111111 + index), None])
    sheet["C3"] = value
    book.save(path); book.close()
    before = path.read_bytes()
    data = parse_dataset(path, path.name)
    profile = profile_dataset(data)
    assert data.columns[-1].display_name == "Kolom 3"
    assert profile.columns[-1].candidate_role == "unclassified"
    assert apply_semantic_roles(profile, {"Kolom 3": "date"}).columns[-1].candidate_role == "unclassified"
    reference = load_administrative_reference(Path(__file__).parents[1] / "reference/master.csv")
    executed = []
    findings = quality_findings(data, profile, reference, executed_checks=executed)
    assert Counter(f.issue_type for f in findings) == {"MISSING_COLUMN_HEADER": 1}
    assert findings[0].row == 3
    assert findings[0].original_value == value
    assert "COLUMN_STRUCTURE" in executed
    assert data.source_rows == list(range(2, 10))
    assert path.read_bytes() == before


def test_csv_prunes_empty_anonymous_columns_and_retains_nonempty_positions(tmp_path):
    path = tmp_path / "anonymous.csv"
    path.write_text('nama,,,catatan\nANI,,EXTRA,N/A\nBUDI,,,N/A\n', encoding="utf-8")
    data = parse_dataset(path, path.name)
    assert [c.display_name for c in data.columns] == ["nama", "Kolom 3", "catatan"]
    findings = structural_findings(data)
    assert [(f.column, f.issue_type) for f in findings] == [("Kolom 3", "MISSING_COLUMN_HEADER")]


def test_anonymous_formula_with_empty_cached_result_is_not_discarded(tmp_path):
    path = tmp_path / "empty-cache.xlsx"
    book = Workbook(); sheet = book.active
    sheet.append(["nama", None]); sheet.append(["ANI", '=""'])
    book.save(path); book.close()
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path, "w") as archive:
        for info, content in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                cell = root.find(".//s:c[@r='B2']", namespace)
                cell.set("t", "str")
                cell.find("s:v", namespace).text = ""
                content = ET.tostring(root)
            archive.writestr(info, content)
    data = parse_dataset(path, path.name)
    assert [column.display_name for column in data.columns] == ["nama", "Kolom 2"]
    assert data.records[0]["Kolom 2"] == '=""'
    assert [f.issue_type for f in structural_findings(data)] == ["MISSING_COLUMN_HEADER"]
