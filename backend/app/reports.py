import csv
import io
import json
from collections import Counter
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Dataset, Finding


DETAIL_HEADERS = [
    "Baris", "Kolom", "Temuan", "Severity", "Nilai saat ini", "Nilai rekomendasi",
    "Evidence", "Rekomendasi tindakan", "Confidence", "Detector",
]


def _value(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def detail_rows(findings: list[Finding]) -> list[list[object]]:
    return [DETAIL_HEADERS] + [[
        item.row,
        item.column or "",
        item.issue_type,
        item.severity,
        _value(item.original_value),
        _value(item.expected_value) if item.expected_value is not None else "",
        item.evidence,
        item.recommendation,
        item.confidence,
        item.detector,
    ] for item in findings]


def summary_rows(dataset: Dataset, findings: list[Finding]) -> list[list[object]]:
    severities = Counter(item.severity for item in findings)
    issues = Counter(item.issue_type for item in findings)
    rows: list[list[object]] = [
        ["Ringkasan assessment", ""],
        ["Dataset", dataset.original_filename],
        ["ID assessment", str(dataset.id)],
        ["Sumber", dataset.source_type],
        ["Worksheet", dataset.worksheet_name or "-"],
        ["Baris", dataset.row_count],
        ["Kolom", dataset.column_count],
        ["Total temuan dalam laporan", len(findings)],
        ["Severity tinggi", severities.get("HIGH", 0)],
        ["Severity medium", severities.get("MEDIUM", 0)],
        ["Selesai diproses", dataset.completed_at.isoformat() if dataset.completed_at else "-"],
        [],
        ["Jenis temuan", "Jumlah"],
    ]
    rows.extend([[issue, count] for issue, count in sorted(issues.items())])
    return rows


def csv_bytes(dataset: Dataset, findings: list[Finding]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerows(summary_rows(dataset, findings))
    writer.writerow([])
    writer.writerows(detail_rows(findings))
    return output.getvalue().encode("utf-8-sig")


def xlsx_bytes(dataset: Dataset, findings: list[Finding]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Ringkasan"
    for row in summary_rows(dataset, findings):
        summary.append(row)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 48
    summary["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    summary["A1"].fill = PatternFill("solid", fgColor="1A73E8")
    summary["B1"].fill = PatternFill("solid", fgColor="1A73E8")

    details = workbook.create_sheet("Temuan")
    for row in detail_rows(findings):
        details.append(row)
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    widths = [10, 22, 27, 13, 32, 32, 48, 48, 13, 28]
    for index, width in enumerate(widths, start=1):
        details.column_dimensions[get_column_letter(index)].width = width
    for cell in details[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A73E8")
    for row in details.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_bytes(dataset: Dataset, findings: list[Finding]) -> bytes:
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=15 * mm,
        bottomMargin=14 * mm,
        title=f"Laporan assessment - {dataset.original_filename}",
        author="Intelligent Data Quality Assessment",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Meta", parent=styles["BodyText"], fontSize=8.5, leading=12, textColor=colors.HexColor("#3C4043")))
    styles.add(ParagraphStyle(name="Cell", parent=styles["BodyText"], fontSize=6.8, leading=8.6, textColor=colors.HexColor("#3C4043")))
    styles.add(ParagraphStyle(name="CellRight", parent=styles["Cell"], alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name="CellHeader", parent=styles["Cell"], textColor=colors.white, fontName="Helvetica-Bold"))
    story = [
        Paragraph("Laporan assessment kualitas data", styles["Title"]),
        Paragraph(escape(dataset.original_filename), styles["Heading2"]),
        Spacer(1, 5 * mm),
    ]
    summary_data = [
        [Paragraph(escape(str(row[0])), styles["Meta"]), Paragraph(escape(str(row[1])), styles["Meta"])]
        for row in summary_rows(dataset, findings)
        if len(row) == 2 and row[0]
    ]
    summary_table = Table(summary_data, colWidths=[58 * mm, 125 * mm], hAlign="LEFT")
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F0FE")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#174EA6")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DADCE0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([summary_table, PageBreak(), Paragraph("Daftar temuan", styles["Heading1"]), Spacer(1, 3 * mm)])
    headers = ["Baris", "Kolom", "Temuan", "Nilai saat ini", "Nilai rekomendasi", "Evidence"]
    table_data = [[Paragraph(header, styles["CellHeader"]) for header in headers]]
    for item in findings:
        table_data.append([
            Paragraph(str(item.row), styles["CellRight"]),
            Paragraph(escape(item.column or "-"), styles["Cell"]),
            Paragraph(f"<b>{escape(item.issue_type)}</b><br/>{escape(item.severity)}", styles["Cell"]),
            Paragraph(escape(_value(item.original_value)), styles["Cell"]),
            Paragraph(escape(_value(item.expected_value)) if item.expected_value is not None else "-", styles["Cell"]),
            Paragraph(escape(item.evidence), styles["Cell"]),
        ])
    if not findings:
        table_data.append([Paragraph("Tidak ada temuan pada filter laporan ini.", styles["Cell"]), "", "", "", "", ""])
    findings_table = Table(table_data, repeatRows=1, colWidths=[13 * mm, 31 * mm, 43 * mm, 47 * mm, 47 * mm, 72 * mm])
    findings_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A73E8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFD")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DADCE0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("SPAN", (0, 1), (-1, 1)) if not findings else ("VALIGN", (0, 1), (-1, -1), "TOP"),
    ]))
    story.append(findings_table)

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#5F6368"))
        canvas.drawString(14 * mm, 8 * mm, "Sumber data tidak diubah. Rekomendasi perlu diverifikasi reviewer.")
        canvas.drawRightString(landscape(A4)[0] - 14 * mm, 8 * mm, f"Halaman {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
