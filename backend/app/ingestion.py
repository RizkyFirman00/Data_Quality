import csv
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


class DatasetParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedColumn:
    original_name: str
    display_name: str
    key: str


@dataclass(frozen=True)
class ParsedDataset:
    columns: list[ParsedColumn]
    records: list[dict[str, object]]
    source_rows: list[int]
    worksheet_name: str | None = None


def _columns(headers: list[object], max_columns: int) -> list[ParsedColumn]:
    if len(headers) > max_columns:
        raise DatasetParseError(f"Dataset melebihi batas {max_columns} kolom.")
    names = ["" if value is None else str(value) for value in headers]
    nonempty = [name for name in names if name]
    if len(nonempty) != len(set(nonempty)):
        raise DatasetParseError("Nama kolom duplikat tidak didukung.")
    columns = [
        ParsedColumn(name, name or f"Kolom {index + 1}", name or f"Kolom {index + 1}")
        for index, name in enumerate(names)
    ]
    keys = [column.key for column in columns]
    if any(len(key) > 255 for key in keys):
        raise DatasetParseError("Nama kolom efektif tidak boleh melebihi 255 karakter.")
    if "_source_row" in keys:
        raise DatasetParseError("Nama kolom _source_row dicadangkan untuk provenance.")
    if len(keys) != len(set(keys)):
        raise DatasetParseError("Nama kolom efektif ambigu setelah header kosong diberi nama.")
    return columns


def _decode_csv(path: Path) -> str:
    raw = path.read_bytes()
    if not raw:
        raise DatasetParseError("File CSV kosong.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise DatasetParseError("Encoding CSV tidak dapat dibaca. Gunakan UTF-8 atau Windows-1252.")


def _parse_csv(
    path: Path, max_rows: int, max_columns: int, max_cells: int
) -> ParsedDataset:
    text = _decode_csv(path)
    if len(text.splitlines()) < 2:
        raise DatasetParseError("CSV harus memiliki header dan minimal satu baris data.")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error as exc:
        if any(delimiter in text.splitlines()[0] for delimiter in ",;\t|"):
            raise DatasetParseError("Delimiter CSV tidak dapat dikenali.") from exc
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    header: list[str] | None = None
    columns: list[ParsedColumn] = []
    records: list[dict[str, object]] = []
    source_rows: list[int] = []
    while True:
        row_number = reader.line_num + 1
        try:
            row = next(reader)
        except StopIteration:
            break
        if not row:
            continue
        if header is None:
            header = row
            if not any(header):
                break
            columns = _columns(header, max_columns)
            if len(columns) > max_cells:
                raise DatasetParseError(f"Dataset melebihi batas {max_cells} sel.")
            continue
        if len(records) >= max_rows:
            raise DatasetParseError(f"Dataset melebihi batas {max_rows} baris data.")
        if (len(records) + 2) * len(columns) > max_cells:
            raise DatasetParseError(f"Dataset melebihi batas {max_cells} sel.")
        if len(row) != len(columns):
            raise DatasetParseError(f"Jumlah kolom tidak konsisten pada baris {row_number}.")
        records.append({column.key: value for column, value in zip(columns, row, strict=True)})
        source_rows.append(row_number)
    if header is None or not any(header):
        raise DatasetParseError("CSV harus memiliki header dan minimal satu baris data.")
    if not records:
        raise DatasetParseError("CSV tidak memiliki baris data.")
    return ParsedDataset(columns, records, source_rows)


def _parse_xlsx(
    path: Path,
    max_uncompressed_bytes: int,
    max_rows: int,
    max_columns: int,
    max_cells: int,
) -> ParsedDataset:
    if not zipfile.is_zipfile(path):
        raise DatasetParseError("File XLSX tidak valid atau rusak.")
    with zipfile.ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > max_uncompressed_bytes:
            raise DatasetParseError("Isi XLSX melebihi batas keamanan.")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        formula_workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        if "workbook" in locals():
            workbook.close()
        raise DatasetParseError("File XLSX tidak dapat dibaca.") from exc
    try:
        for sheet in workbook.worksheets:
            formula_sheet = formula_workbook[sheet.title]
            if sheet.max_column > max_columns:
                raise DatasetParseError(f"Dataset melebihi batas {max_columns} kolom.")
            if sheet.max_row * sheet.max_column > max_cells:
                raise DatasetParseError(f"Dataset melebihi batas {max_cells} sel.")
            columns: list[ParsedColumn] | None = None
            records: list[dict[str, object]] = []
            source_rows: list[int] = []
            rows = zip(
                sheet.iter_rows(values_only=True),
                formula_sheet.iter_rows(),
                strict=True,
            )
            for row_number, (row, formula_row) in enumerate(rows, start=1):
                if columns is None:
                    if not any(value is not None for value in row):
                        continue
                    columns = _columns(list(row), max_columns)
                    continue
                has_formula = any(cell.data_type == "f" for cell in formula_row)
                if not any(value is not None for value in row) and not has_formula:
                    continue
                if len(records) >= max_rows:
                    raise DatasetParseError(f"Dataset melebihi batas {max_rows} baris data.")
                records.append(
                    {
                        column.key: (
                            value
                            if value is not None or formula_cell.data_type != "f"
                            else (
                                formula_cell.value
                                if isinstance(formula_cell.value, str)
                                else getattr(
                                    formula_cell.value,
                                    "text",
                                    str(formula_cell.value),
                                )
                            )
                        )
                        for column, value, formula_cell in zip(
                            columns, row, formula_row, strict=True
                        )
                    }
                )
                source_rows.append(row_number)
            if columns is None:
                continue
            if not records:
                raise DatasetParseError("Worksheet harus memiliki minimal satu baris data.")
            return ParsedDataset(columns, records, source_rows, sheet.title)
    finally:
        workbook.close()
        formula_workbook.close()
    raise DatasetParseError("Workbook tidak memiliki worksheet berisi data.")


def parse_dataset(
    path: Path,
    original_name: str,
    max_uncompressed_bytes: int = 250 * 1024 * 1024,
    max_rows: int = 100_000,
    max_columns: int = 1_000,
    max_cells: int = 2_000_000,
) -> ParsedDataset:
    suffix = Path(original_name).suffix.lower()
    if suffix == ".csv":
        return _parse_csv(path, max_rows, max_columns, max_cells)
    if suffix == ".xlsx":
        return _parse_xlsx(
            path, max_uncompressed_bytes, max_rows, max_columns, max_cells
        )
    raise DatasetParseError("Format file tidak didukung. Gunakan CSV atau XLSX.")
